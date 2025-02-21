from openpyxl import load_workbook
from sqlalchemy import and_
from tqdm import tqdm
import os

import db_models
import database
from sche_che import FuncToolBox

DATA_FOLDER = 'data'
SCHEDULE_FOLDER = 'current_version_schedules'
SUBJECTS = ['английский язык', 'биология', 'география', 'ИЗО', 'информатика', 'история', 'литература', 'математика',
            'музыка', 'ОБЗР', 'обществознание', 'ОДКНР', 'природоведение', 'проект', 'профориентация', 'РоВ',
            'русский язык', 'технология', 'физ. час', 'физика', 'физкультура', 'химия']


def create_subjects(subj_list):
    with db:
        for subj in subj_list:
            subject = db_models.Subject(name=subj)
            db.add(subject)
        db.commit()
    print('subject created!')


def create_classes_from_file(filename):
    wb = load_workbook(filename)
    ws = wb.active
    start_row = 3
    with db:
        for i in range(start_row, ws.max_row + 1):
            new_class = db_models.Class(name=ws.cell(i, 2).value, shift=int(ws.cell(i, 3).value),
                                        quantity=int(ws.cell(i, 4).value), lessons_min=int(ws.cell(i, 5).value),
                                        lessons_max=int(ws.cell(i, 6).value))
            db.add(new_class)
        db.commit()
    print('classes created!')


def create_cabinets_from_file(filename):
    wb = load_workbook(filename)
    ws = wb.active
    start_row = 3
    with db:
        for i in range(start_row, ws.max_row + 1):
            cabinet = db_models.Cabinet(number=ws.cell(i, 2).value, capacity=ws.cell(i, 3).value)
            db.add(cabinet)
        db.commit()
    print('cabinets created!')


def create_teachers_from_file(filename):
    wb = load_workbook(filename)
    ws = wb.active
    start_row = 3
    with db:
        for i in range(start_row, ws.max_row + 1):
            cabinet_number = ws.cell(i, 5).value
            cabinet_id = db.session.query(db_models.Cabinet.id).filter(db_models.Cabinet.number == cabinet_number).one()[0]
            teacher = db_models.Teacher(surname=ws.cell(i, 2).value, name_last_name=ws.cell(i, 3).value,
                                        base_cabinet=cabinet_id)
            db.add(teacher)
        db.commit()
    print('teachers created!')


def create_teachers_specializations(filename):
    wb = load_workbook(filename)
    ws = wb.active
    start_row = 3
    with db:
        for i in range(start_row, ws.max_row + 1):
            current_surname = ws.cell(i, 2).value
            name_last_name = ws.cell(i, 3).value
            teacher_id = db.session.query(db_models.
                                          Teacher.id).filter(and_(db_models.Teacher.surname == current_surname,
                                                                  db_models.Teacher.name_last_name == name_last_name)).one()[0]
            subjects = ws.cell(i, 4).value.split(', ')
            for subj in subjects:
                subject_id = db.session.query(db_models.Subject.id).filter(db_models.Subject.name == subj).one()[0]
                specialization = db_models.TeacherSpecialization(teacher_id=teacher_id, subject_id=subject_id)
                db.add(specialization)
        db.commit()
    print('specializations created!')


def create_class_teacher_subject_connetions(filename):
    wb = load_workbook(filename)
    ws = wb.active
    start_row = 3
    with db:
        for i in range(start_row, ws.max_row):
            class_name = ws.cell(i, 2).value
            hours = ws.cell(i, 3).value
            if not class_name:
                class_name = ws.cell(i - 1, 2).value
                hours = ws.cell(i - 1, 3).value
            class_name_list = class_name.split(', ')
            hours = int(hours)
            subject = ws.cell(i, 6).value
            teacher = ws.cell(i, 7).value
            teacher_surname = teacher.split(' ')[0]
            cabinet = ws.cell(i, 8).value
            week = 'odd' if ws.cell(i, 11).value == 'Неч' else 'even'

            subject_id = db.session.query(db_models.Subject.id).filter(db_models.Subject.name == subject).one()[0]
            _teachers = db.session.query(db_models.Teacher).filter(db_models.Teacher.surname == teacher_surname).all()
            teacher_id = None
            if len(_teachers) > 1:
                for _t in _teachers:
                    x = _t.name_last_name.split(' ')[1][0]
                    if x == teacher.split(' ')[2][0]:
                        teacher_id = _t.id
                        break
            else:
                teacher_id = _teachers[0].id
            for class_name in class_name_list:
                class_id = db.session.query(db_models.Class.id).filter(db_models.Class.name == class_name).one()[0]
                if cabinet:
                    cabinet_id = \
                    db.session.query(db_models.Cabinet.id).filter(db_models.Cabinet.number == cabinet).one()[0]
                    connection = db_models.ClassTeacherSubjectConnection(class_id=class_id, times_a_week=hours,
                                                                     subject_id=subject_id, teacher_id=teacher_id,
                                                                     cabinet_id=cabinet_id)
                else:
                    connection = db_models.ClassTeacherSubjectConnection(class_id=class_id, times_a_week=hours,
                                                                         subject_id=subject_id, teacher_id=teacher_id)
                db.add(connection)
        db.commit()
    print('connections created!')


def create_schedule_from_file(filename):
    # что делать с уроками по чет-нечет? пока рассматривать как обычные
    # замораживаем функцию: неясно по каким данным определять какой конкретно учиель англ в каком кабинете
    # они не по алфавиту. придется таки смотреть в выгрузку расписания учителей
    wb = load_workbook(filename)
    toolbox = FuncToolBox()
    if 'NORM' not in filename:
        wb = toolbox.row_normalization(wb)
    ws = wb.active
    row = 1
    with db:
        while row < ws.max_row:
            if 'Класс' in (x := str(ws.cell(row, 1).value)):
                this_class = x.split(' - ')[1]
                class_id = db.session.query(db_models.Class.id).filter(db_models.Class.name == this_class).one()[0]
                print(this_class)
                row += 3
                while lesson_n := str(ws.cell(row, 1).value):
                    if lesson_n == 'None':
                        break
                    lesson_n = int(lesson_n.split(':')[-1])
                    for day in range(1, 6):
                        subject = str(ws.cell(row, day * 2).value)
                        cabinet = str(ws.cell(row, day * 2 + 1).value)
                        if subject == 'None':
                            continue
                        subjects = [s.replace('(н)', '') for s in subject.split('\n')]
                        cabinets = [c.replace('(н)', '') for c in cabinet.split('\n')]
                        subject_cabinet_pairs = []
                        if len(cabinets) > len(subjects):
                            subject_cabinet_pairs.extend([(subjects[0], cabinets[0]), (subjects[0], cabinets[1])])
                        else:
                            for s in subjects:
                                for c in cabinets:
                                    subject_cabinet_pairs.append((s, c))
                        for s, c in subject_cabinet_pairs:

                            subject_id = db.session.query(db_models.Subject.id).filter(db_models.Subject.name == s).one()[0]
                            cabinet_id = db.session.query(db_models.Cabinet.id).filter(db_models.Cabinet.number == c).one()[0]
                            teacher_id = db.session.query(
                                db_models.ClassTeacherSubjectConnection.teacher_id).with_entities(
                                db_models.ClassTeacherSubjectConnection.teacher_id).filter(
                                db_models.ClassTeacherSubjectConnection.subject_id == subject_id,
                                db_models.ClassTeacherSubjectConnection.class_id == class_id).all()
                            if s == 'английский язык':
                                print(f'{day}, {lesson_n}')
                                print(s, c)
                                print(teacher_id)
                                for t in teacher_id:
                                    print(db.session.query(db_models.Teacher.surname).filter(db_models.Teacher.id == t[0]).one()[0], end= ' ')
                                print()
                    row += 1
            row += 1



if __name__ == '__main__':
    db = database.Database()
    db.init_db()
    # db.recreate_db()
    # create_subjects(SUBJECTS)
    # create_classes_from_file(os.path.join(DATA_FOLDER, 'classes_list.xlsx'))
    # create_cabinets_from_file(os.path.join(DATA_FOLDER, 'rooms_list.xlsx'))
    # create_teachers_from_file(os.path.join(DATA_FOLDER, 'teachers_list.xlsx'))
    # create_teachers_specializations(os.path.join(DATA_FOLDER, 'teachers_list.xlsx'))
    # create_class_teacher_subject_connetions(os.path.join(DATA_FOLDER, 'connections.xlsx'))
    create_schedule_from_file(os.path.join(SCHEDULE_FOLDER, 'расписание учеников_NORM.xlsx'))


